from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(r"C:\Users\18712\Desktop\SmartHome")
SRC = ROOT / "专业实习报告-SmartHome智能家居控制系统-图文版.docx"
OUT = ROOT / "专业实习报告-SmartHome智能家居控制系统-图文增强版.docx"
WORK = ROOT / "report_assets"


def get_font(size=28, bold=False):
    candidates = [
        r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\simsun.ttc",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def center_text(draw, box, text, fill, size=24, bold=False):
    font = get_font(size, bold)
    lines = text.split("\n")
    line_heights = []
    line_widths = []
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        line_widths.append(bbox[2] - bbox[0])
        line_heights.append(bbox[3] - bbox[1] + 8)
    total_h = sum(line_heights)
    y = box[1] + (box[3] - box[1] - total_h) / 2
    for idx, line in enumerate(lines):
        x = box[0] + (box[2] - box[0] - line_widths[idx]) / 2
        draw.text((x, y), line, fill=fill, font=font)
        y += line_heights[idx]


def arrow(draw, start, end, color="#64748B", width=4):
    draw.line((start[0], start[1], end[0], end[1]), fill=color, width=width)
    dx = end[0] - start[0]
    dy = end[1] - start[1]
    if abs(dx) >= abs(dy):
        if dx >= 0:
            pts = [(end[0], end[1]), (end[0] - 13, end[1] - 8), (end[0] - 13, end[1] + 8)]
        else:
            pts = [(end[0], end[1]), (end[0] + 13, end[1] - 8), (end[0] + 13, end[1] + 8)]
    else:
        if dy >= 0:
            pts = [(end[0], end[1]), (end[0] - 8, end[1] - 13), (end[0] + 8, end[1] - 13)]
        else:
            pts = [(end[0], end[1]), (end[0] - 8, end[1] + 13), (end[0] + 8, end[1] + 13)]
    draw.polygon(pts, fill=color)


def rounded_box(draw, box, title, body="", fill="#FFFFFF", outline="#CBD5E1", accent="#2563EB"):
    draw.rounded_rectangle(box, radius=18, fill=fill, outline=outline, width=2)
    if body:
        draw.rectangle((box[0], box[1], box[2], box[1] + 42), fill=accent)
        draw.rounded_rectangle((box[0], box[1], box[2], box[3]), radius=18, outline=outline, width=2)
        center_text(draw, (box[0], box[1] + 2, box[2], box[1] + 40), title, "#FFFFFF", 21, True)
        center_text(draw, (box[0] + 8, box[1] + 46, box[2] - 8, box[3] - 8), body, "#374151", 18)
    else:
        center_text(draw, box, title, "#111827", 22, True)


def make_function_structure():
    out = Image.new("RGB", (1180, 780), "#F6F7F9")
    draw = ImageDraw.Draw(out)
    draw.text((40, 28), "SmartHome 功能结构图", fill="#111827", font=get_font(34, True))
    center = (440, 108, 740, 188)
    rounded_box(draw, center, "SmartHome 智能家居控制系统", fill="#EEF6FF", outline="#93C5FD")
    nodes = [
        ((60, 300, 300, 410), "首页中控", "空间总览\n日夜背景\n场景快捷入口"),
        ((350, 300, 590, 410), "房间控制", "客厅/卧室切换\n点位坐标\n设备状态"),
        ((640, 300, 880, 410), "设备详情", "指标卡片\n滑块调节\n托管/开关"),
        ((930, 300, 1170, 410), "智能助手", "自然语言输入\n本地规则解析\nDeepSeek 预留"),
        ((210, 560, 450, 670), "自动化脚本", "触发条件\n动作列表\n启用状态"),
        ((730, 560, 970, 670), "资源与模型", "房间状态图\n设备图片\nGLB/GLTF 模型"),
    ]
    for box, title, body in nodes:
        rounded_box(draw, box, title, body, accent="#334155")
        arrow(draw, ((center[0] + center[2]) // 2, center[3]), ((box[0] + box[2]) // 2, box[1] - 4))
    out_path = WORK / "function_structure.png"
    out.save(out_path, quality=94)
    return out_path


def make_technical_principle():
    out = Image.new("RGB", (1180, 760), "#F6F7F9")
    draw = ImageDraw.Draw(out)
    draw.text((40, 28), "SmartHome 技术原理图", fill="#111827", font=get_font(34, True))
    layers = [
        ((80, 110, 1100, 210), "表现层 ArkTS UI", "Index.ets、DeviceDetailPages.ets；负责页面布局、底部导航、弹层、滑块、卡片和图文展示", "#2563EB"),
        ((80, 250, 1100, 350), "状态层 State / Prop / Link", "activeRoomId、selectedPointId、isNight、设备开关、温度、窗帘比例等状态驱动界面刷新", "#0F766E"),
        ((80, 390, 1100, 490), "业务层 房间与设备模型", "RoomSpec、PointSpec、DevicePageSpec、AutomationSpec；把家庭空间、设备点位和脚本抽象成结构化数据", "#7C3AED"),
        ((80, 530, 1100, 630), "智能解析与资源层", "SmartPromptParser 解析指令；media/rawfile 提供房间图、设备图、3D 模型；DeepSeekService 作为普通问答扩展", "#B45309"),
    ]
    for idx, (box, title, body, accent) in enumerate(layers):
        rounded_box(draw, box, title, body, fill="#FFFFFF", outline="#D1D5DB", accent=accent)
        if idx < len(layers) - 1:
            arrow(draw, ((box[0] + box[2]) // 2, box[3] + 8), ((box[0] + box[2]) // 2, layers[idx + 1][0][1] - 8))
    out_path = WORK / "technical_principle.png"
    out.save(out_path, quality=94)
    return out_path


def make_algorithm_flow():
    out = Image.new("RGB", (1180, 840), "#F6F7F9")
    draw = ImageDraw.Draw(out)
    draw.text((40, 28), "智能助手自动化生成算法流程图", fill="#111827", font=get_font(34, True))
    boxes = {
        "start": (450, 95, 730, 155),
        "input": (410, 205, 770, 275),
        "judge": (425, 330, 755, 430),
        "trigger": (90, 500, 390, 590),
        "actions": (450, 500, 750, 590),
        "llm": (810, 500, 1110, 590),
        "save": (265, 690, 565, 775),
        "reply": (625, 690, 925, 775),
    }
    rounded_box(draw, boxes["start"], "开始")
    rounded_box(draw, boxes["input"], "输入用户指令\nagentPrompt")
    diamond = [(590, 315), (770, 380), (590, 445), (410, 380)]
    draw.polygon(diamond, fill="#FFFFFF", outline="#CBD5E1")
    center_text(draw, (430, 332, 750, 428), "是否包含\n家居控制关键词", "#111827", 22, True)
    rounded_box(draw, boxes["trigger"], "解析触发条件", "每天/每周\n离家/回家\n手动执行", accent="#0F766E")
    rounded_box(draw, boxes["actions"], "解析设备动作", "灯光/空调/窗帘\n电视/净化器\n门锁/扫地机器人", accent="#7C3AED")
    rounded_box(draw, boxes["llm"], "普通问答分支", "调用 DeepSeekService\n或提示配置 API Key", accent="#64748B")
    rounded_box(draw, boxes["save"], "生成 AutomationSpec", "title、trigger\ntriggerSpec、actions\nenabled", accent="#B45309")
    rounded_box(draw, boxes["reply"], "更新界面反馈", "聊天消息\n草稿卡片\n脚本列表", accent="#2563EB")
    arrow(draw, (590, 155), (590, 205))
    arrow(draw, (590, 275), (590, 315))
    arrow(draw, (410, 380), (240, 500))
    arrow(draw, (590, 445), (590, 500))
    arrow(draw, (770, 380), (960, 500))
    arrow(draw, (240, 590), (415, 690))
    arrow(draw, (590, 590), (415, 690))
    arrow(draw, (960, 590), (775, 690))
    arrow(draw, (565, 735), (625, 735))
    draw.text((300, 452), "是", fill="#0F766E", font=get_font(20, True))
    draw.text((790, 452), "否", fill="#64748B", font=get_font(20, True))
    out_path = WORK / "algorithm_flow.png"
    out.save(out_path, quality=94)
    return out_path


def set_caption(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(8)
    for run in paragraph.runs:
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.size = Pt(10.5)


def insert_before(paragraph, image_path, caption, width=6.0):
    image_p = paragraph.insert_paragraph_before()
    image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_p.add_run().add_picture(str(image_path), width=Inches(width))
    cap_p = paragraph.insert_paragraph_before(caption)
    set_caption(cap_p)


def insert_before_contains(doc, marker, image_path, caption, width=6.0):
    for paragraph in doc.paragraphs:
        if marker in paragraph.text:
            insert_before(paragraph, image_path, caption, width)
            return True
    return False


def main():
    WORK.mkdir(exist_ok=True)
    function_structure = make_function_structure()
    technical_principle = make_technical_principle()
    algorithm_flow = make_algorithm_flow()

    doc = Document(SRC)
    inserted = []
    inserted.append(insert_before_contains(
        doc,
        "系统功能结构可概括为",
        function_structure,
        "图4-8 SmartHome 功能结构图",
        6.1,
    ))
    inserted.append(insert_before_contains(
        doc,
        "系统采用单 entry 模块组织",
        technical_principle,
        "图4-9 SmartHome 技术原理图",
        6.1,
    ))
    inserted.append(insert_before_contains(
        doc,
        "第三，智能助手采用本地规则优先的策略",
        algorithm_flow,
        "图4-10 智能助手自动化生成算法流程图",
        6.1,
    ))
    doc.save(OUT)
    print(OUT)
    print(f"inserted={sum(1 for item in inserted if item)}")
    print(f"images={len(doc.inline_shapes)} paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}")


if __name__ == "__main__":
    main()
