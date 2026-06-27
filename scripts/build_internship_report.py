from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\18712\Desktop\SmartHome")
OUT = ROOT / "专业实习报告-SmartHome智能家居控制系统.docx"
MEDIA = ROOT / "entry" / "src" / "main" / "resources" / "base" / "media"


def set_run_font(run, font="宋体", size=None, bold=None, color=None):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_para(paragraph, font="宋体", size=12, line=1.5, first_line=True, align=None, before=0, after=6):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if first_line:
        fmt.first_line_indent = Cm(0.74)
    if align is not None:
        paragraph.alignment = align
    for run in paragraph.runs:
        set_run_font(run, font, size)


def add_text(doc, text, first_line=True, after=6):
    p = doc.add_paragraph()
    p.add_run(text)
    set_para(p, first_line=first_line, after=after)
    return p


def add_heading_cn(doc, text, level=1):
    p = doc.add_paragraph()
    r = p.add_run(text)
    if level == 1:
        set_run_font(r, "黑体", 15, True)
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(8)
    elif level == 2:
        set_run_font(r, "黑体", 13, True)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(6)
    else:
        set_run_font(r, "宋体", 12, True)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.5
    return p


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    r = p.add_run(text)
    set_run_font(r, "宋体", 11, bold)
    p.paragraph_format.space_after = Pt(0)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, True)
        shade_cell(table.rows[0].cells[i], "E8EEF5")
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], val, False, WD_ALIGN_PARAGRAPH.LEFT if i > 0 else WD_ALIGN_PARAGRAPH.CENTER)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)
    doc.add_paragraph()
    return table


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run_font(r, "宋体", 10.5)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)


def add_picture(doc, path, caption, width=5.6):
    if not Path(path).exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


def add_code(doc, code):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.3)
    p.paragraph_format.right_indent = Cm(0.2)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.05
    for line in code.strip().splitlines():
        r = p.add_run(line.rstrip() + "\n")
        set_run_font(r, "Consolas", 8.5)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.4)

    styles = doc.styles
    styles["Normal"].font.name = "宋体"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    styles["Normal"].font.size = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(30)
    r = p.add_run("专业实习报告")
    set_run_font(r, "黑体", 24, True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("（SmartHome 智能家居控制系统）")
    set_run_font(r, "黑体", 16, True)

    cover = doc.add_table(rows=5, cols=2)
    cover.alignment = WD_TABLE_ALIGNMENT.CENTER
    cover.style = "Table Grid"
    cover_data = [
        ("专业班级：", "计UC-23x"),
        ("学    号：", ""),
        ("姓    名：", ""),
        ("指导老师：", "熊静、田延安、姚斌、谢佳、贾小云、赵晓、李婉、王梅嘉"),
        ("实习地点：", "软通动力"),
    ]
    for idx, (label, value) in enumerate(cover_data):
        set_cell_text(cover.rows[idx].cells[0], label, True, WD_ALIGN_PARAGRAPH.RIGHT)
        set_cell_text(cover.rows[idx].cells[1], value, False, WD_ALIGN_PARAGRAPH.LEFT)
        cover.rows[idx].cells[0].width = Cm(4)
        cover.rows[idx].cells[1].width = Cm(10)

    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("实习时间：2026年 6 月 1 日至 2026年 7 月 3 日")
    set_run_font(r, "宋体", 12, True)

    doc.add_page_break()

    add_heading_cn(doc, "一、专业实习目的及意义")
    add_text(doc, "这次专业实习对我来说不是单纯把课堂案例重新敲一遍，而是把前面分散学到的鸿蒙页面、资源管理、Git、Linux 和 OpenHarmony 基础知识放到一个连续项目里验证。我的结课作业选择了 SmartHome 智能家居控制系统，原因是这个题目既有可视化界面，又有设备状态、场景联动和智能助手逻辑，能够暴露出“页面好看”和“系统能维护”之间的差别。")
    add_text(doc, "在项目推进过程中，我逐渐意识到智能家居应用的难点并不只是摆几个设备按钮，而是怎样让用户一眼看懂家里发生了什么。例如客厅灯光打开后，房间背景图要跟着变化；进入卧室后，点位和设备详情不能继续显示客厅的数据；智能助手解析自然语言后，也要把文字变成明确的触发条件和设备动作。这些问题促使我从状态建模、组件拆分、资源命名和交互反馈几个角度重新组织代码。")
    add_text(doc, "因此，本次实习的主要意义在于让我把“能实现一个效果”的思路推进到“能解释一个系统为什么这样设计”。我不再只关注某个按钮是否能点，而是开始考虑数据结构是否方便扩展、页面组件是否复用、异常情况是否有提示、第三方资源是否保留许可说明。这个过程让我更接近企业项目中的工程思维。")

    add_heading_cn(doc, "二、实习单位简介")
    add_text(doc, "本次实习由学校与软通动力相关实训资源共同组织完成。对我个人而言，软通动力这个实习环境最有价值的地方，不只是介绍一家企业的规模和业务，而是把企业研发中常见的工作方式带进了课程：先明确需求，再拆模块；先搭出可运行版本，再逐步补齐交互和异常；遇到陌生工具时，先按文档完成最小可用流程，再继续深入。")
    add_text(doc, "实习内容与鸿蒙生态和 OpenHarmony 场景结合比较紧密。课程中涉及移动端页面开发、设备侧开发环境、Git 版本管理、Linux 命令、Makefile 构建思想等内容，这些知识在 SmartHome 项目中并不是孤立出现的。例如资源图片和模型文件需要按项目目录组织，页面逻辑需要用组件和状态管理承接，智能家居设备的抽象又需要参考物联网场景中的“设备、状态、动作、触发条件”模型。")
    add_text(doc, "相比只听企业介绍，我更大的收获是开始理解软件服务企业为什么强调规范、交付和协作。一个应用在个人电脑上能运行只是第一步，真正进入企业项目后，还要面对多人维护、需求变化、测试验证、版权合规和用户体验等问题。这些内容直接影响我后面写代码和整理报告时的取舍。")

    add_heading_cn(doc, "三、实习总体安排")
    add_table(
        doc,
        ["项目", "具体内容"],
        [
            ["实习时间", "2026年6月1日至2026年7月3日，共5周。"],
            ["实习地点", "软通动力实训环境及个人开发环境。"],
            ["实习方向", "鸿蒙应用开发与 OpenHarmony 智能终端/智能家居方向。"],
            ["主要内容", "鸿蒙基础、ArkTS 声明式开发、组件与状态管理、Git、Linux、Makefile、RK2206/OpenHarmony 设备侧基础、智能家居综合项目实践。"],
            ["最终成果", "完成 SmartHome 智能家居控制系统，包括首页空间总览、房间点位控制、设备详情页、智能助手与自动化脚本管理等功能。"],
        ],
        [3.0, 12.0],
    )
    add_text(doc, "整体安排上，前期以基础知识学习和案例练习为主，掌握鸿蒙应用项目结构、ArkTS 语法、组件布局、资源管理和开发工具使用；中期学习 Git、Linux、Makefile 与 OpenHarmony/RK2206 设备侧相关知识，理解软件研发与设备开发之间的关系；后期围绕智能家居结课作业进行需求整理、界面设计、功能开发、调试测试和报告总结。")

    add_heading_cn(doc, "四、实习内容及过程")
    add_heading_cn(doc, "4.1 企业项目实践：SmartHome 智能家居控制系统", 2)
    add_text(doc, "SmartHome 是一个面向家庭住户的鸿蒙智能家居控制应用。系统以“先看到家，再进行控制”为设计原则，打开应用后呈现空间化的家庭中控视图，用户可以进入客厅、卧室等房间，在俯视图中查看设备位置、状态和可执行操作，并通过设备详情页完成开关、亮度、温度、窗帘比例、托管模式等交互。同时，系统提供智能助手页面，将自然语言指令解析为设备动作或自动化脚本草稿。")
    add_picture(doc, MEDIA / "home_core_scene.png", "图4-1 SmartHome 首页核心场景资源", 5.2)

    add_heading_cn(doc, "4.1.1 可行性研究", 3)
    add_text(doc, "经济可行性方面，SmartHome 以课程实训和个人结课作业为目标，主要使用 DevEco Studio、HarmonyOS SDK、ArkTS、项目内图片资源以及开源 3D 家具模型资源完成开发，不需要额外购买硬件设备和商业服务。项目中的设备状态采用本地状态模拟，智能助手优先通过本地规则解析，只有普通问答场景才预留 DeepSeek API 调用，因此开发成本和运行成本较低。")
    add_text(doc, "技术可行性方面，鸿蒙 ArkTS 声明式 UI 适合构建卡片、列表、滑块、弹层和底部导航等移动端交互。项目已具备完整工程结构，包括 AppScope、entry 模块、资源文件、页面组件和智能助手服务代码。通过 @State、@Prop、@Link 等状态机制，可以实现房间切换、设备选择、详情页联动和自动化脚本列表更新。资源目录中包含客厅、卧室日夜灯光状态图、设备预览图以及 glb/gltf 模型资源，能够支撑图文并茂的空间化界面。")
    add_text(doc, "社会可行性方面，智能家居是居民生活数字化的重要方向。一个清晰、可访问、可理解的控制界面能够降低普通用户管理多设备的难度，提高节能、安全和生活便利性。项目虽然是课程级原型，但围绕照明、空调、窗帘、电视、空气净化器、扫地机器人和安防门锁等常见家庭设备进行建模，具有较好的应用场景代表性。")

    add_heading_cn(doc, "4.1.2 需求分析", 3)
    add_text(doc, "系统服务对象主要是使用 HarmonyOS 手机、平板或二合一设备的家庭住户。用户希望在一个入口中快速了解家庭当前状态，并直接完成常用设备控制，不希望在复杂菜单中逐层查找设备。因此，系统需求可分为场景感知、房间管理、设备控制、设备详情、智能助手和自动化脚本六类。")
    add_table(
        doc,
        ["用户需求", "功能分解", "实现要点"],
        [
            ["查看家庭状态", "首页展示客厅、卧室等空间入口和设备摘要。", "使用房间数据模型和资源图片，根据日夜状态与灯光状态切换场景图。"],
            ["控制房间设备", "在房间俯视图中点击点位，执行开关、调温、调光、窗帘比例等操作。", "通过 PointSpec 描述点位坐标、设备类型、详情和动作类型。"],
            ["查看设备详情", "进入设备详情页查看指标、设备能力、联动设备和控制滑块。", "DeviceDetailPages 组件根据 pointId 匹配设备页面配置。"],
            ["管理自动化", "创建离家断电、夜间入睡、回家模式、周末清洁等自动化脚本。", "AutomationSpec 保存触发条件、动作列表、来源和启用状态。"],
            ["智能助手交互", "输入自然语言指令，解析为设备动作或脚本草稿。", "SmartPromptParser 负责关键词识别、触发条件解析和动作生成。"],
            ["适配视觉状态", "支持日间/夜间视觉、按压反馈、弹层过渡和清晰状态提示。", "通过颜色函数、动效函数和状态变量统一管理视觉反馈。"],
        ],
        [2.6, 6.0, 6.4],
    )
    add_text(doc, "用例层面，用户进入应用后首先浏览首页中控视图；选择房间后进入房间控制界面；点击设备点位可以打开详情面板；在详情面板中可调整设备参数或启用智能托管；进入智能页面后，用户可以输入“每天22:30关闭客厅灯、拉上卧室窗帘并把卧室空调调到23度”等指令，系统生成自动化草稿，用户确认后保存为脚本。")

    add_heading_cn(doc, "4.1.3 概要设计", 3)
    add_text(doc, "系统采用单 entry 模块组织，核心页面为 Index.ets，设备详情抽象为 DeviceDetailPages.ets，智能助手相关逻辑放在 smart 目录下。整体结构体现了“页面展示层、业务状态层、智能解析层、资源层”的分层思路。页面展示层负责 UI 和交互，业务状态层负责房间、设备和自动化状态，智能解析层负责自然语言到结构化动作的转换，资源层提供图片、模型和配置文件。")
    add_table(
        doc,
        ["模块", "主要文件", "职责"],
        [
            ["入口与主页面", "entry/src/main/ets/pages/Index.ets", "定义房间、点位、场景、导航、弹层、智能页与全局状态。"],
            ["设备详情", "entry/src/main/ets/components/DeviceDetailPages.ets", "根据设备 id 展示设备封面、指标、滑块控制、动作列表和联动设备。"],
            ["智能模型", "entry/src/main/ets/smart/SmartModels.ets", "定义消息、触发器、动作、自动化脚本等接口。"],
            ["指令解析", "entry/src/main/ets/smart/SmartPromptParser.ets", "识别时间、离家、回家、清洁、灯光、空调、窗帘等关键词并生成动作。"],
            ["助手服务", "entry/src/main/ets/smart/SmartAssistantService.ets", "判断本地智能家居指令或转发到 DeepSeek 服务。"],
            ["资源文件", "resources/base/media 与 rawfile/models", "保存房间图、设备图、模型资源和许可说明。"],
        ],
        [2.6, 5.2, 7.2],
    )
    add_text(doc, "系统功能结构可概括为：SmartHome 应用包含首页、房间控制、设备详情、智能助手、自动化脚本和资源模型六个部分。其中首页负责空间入口和整体状态，房间控制负责点位交互，设备详情负责单设备精细控制，智能助手负责自然语言处理，自动化脚本负责触发条件与动作列表的持久化表达。")
    add_picture(doc, MEDIA / "room_living_scene.png", "图4-2 客厅空间场景资源", 5.5)
    add_picture(doc, MEDIA / "room_bedroom_scene.png", "图4-3 卧室空间场景资源", 5.5)

    add_heading_cn(doc, "4.1.4 详细设计", 3)
    add_text(doc, "在数据模型设计上，Index.ets 中定义了 RoomSpec、PointSpec、SceneSpec、RoomDeviceState 等接口。RoomSpec 表示房间基本信息和场景图片集合，PointSpec 表示设备点位，包括所属房间、显示名称、类型、x/y 百分比坐标、详情文本和动作类型。通过点位坐标，设备可以被放置在房间图片中的具体位置，使控制入口具备空间语义。")
    add_text(doc, "在状态设计上，主页面使用多个 @State 变量维护界面状态，例如 activeTab 表示当前底部导航页，activeRoomId 表示当前房间，selectedPointId 和 detailPointId 表示当前选中设备与详情面板，isNight 表示日夜模式，livingLightOn、bedroomLightOn、livingTemperature、bedroomTemperature 等变量表示设备状态。通过这些状态变量，界面能够在用户点击、滑动、输入后自动刷新。")
    add_text(doc, "在页面交互设计上，首页使用底部导航切换“家、智能、商城、我的”等页面；房间页面通过卡片、点位和弹层组织信息；设备详情页使用 Header、Hero、ControlDeck、Metrics、ActionList 等 Builder 拆分页面结构。这样既降低了单个 UI 片段的复杂度，也便于后续增加设备类型。")
    add_text(doc, "在智能助手设计上，SmartPromptParser 通过关键词匹配识别触发条件和动作。触发条件包括每天、每周、离家、回家、手动等类型；动作包括灯光开关、空调温度设置、窗帘比例、电视开关、净化器启动、扫地机器人运行和门锁安防提示等。解析结果被组织为 SmartPromptDraft，并进一步转化为 AutomationSpec 保存到自动化列表。")
    add_table(
        doc,
        ["页面/功能", "关键状态或数据", "交互结果"],
        [
            ["首页中控", "activeTab、rooms、scenes、isNight", "展示空间入口、场景快捷入口和日夜背景。"],
            ["房间点位", "activeRoomId、roomDevices、selectedPointId", "筛选当前房间已安装设备，并显示可点击点位。"],
            ["设备详情", "detailPointId、currentValue、powerOn、autoManaged", "打开详情面板，调整设备参数并反馈指标变化。"],
            ["智能助手", "agentPrompt、chatMessages、agentThinking", "输入指令后生成回复、草稿和自动化动作。"],
            ["自动化脚本", "automations、expandedAutomationId", "展示脚本列表，支持展开查看触发条件和动作。"],
        ],
        [3.0, 5.2, 6.8],
    )

    add_heading_cn(doc, "4.1.5 编码实现", 3)
    add_text(doc, "本项目的编码重点是 ArkTS 声明式 UI、状态驱动交互和智能指令解析。下面选取项目中的核心实现思想进行说明。")
    add_text(doc, "第一，房间图片根据日夜模式和灯光状态动态切换。系统通过当前房间、是否夜间、灯光是否开启三个条件选择对应资源，使用户在视觉上直接感知房间状态。")
    add_code(doc, """
private roomImageById(roomId: string): Resource {
  const room = this.roomById(roomId);
  if (room.images === undefined) {
    return $r('app.media.room_living_day_lightoff');
  }
  const lightOn = this.lightOnForRoom(roomId);
  if (this.isNight) {
    return lightOn ? room.images!.nightLightOn : room.images!.nightLightOff;
  }
  return lightOn ? room.images!.dayLightOn : room.images!.dayLightOff;
}
""")
    add_text(doc, "第二，设备详情组件通过 pointId 获取配置项，避免为每个设备重复编写完整页面。设备标题、图片、指标、控制滑块和动作列表均来自 DevicePageSpec，实现了配置驱动的页面复用。")
    add_code(doc, """
private page(): DevicePageSpec {
  for (let i = 0; i < DEVICE_PAGES.length; i++) {
    if (DEVICE_PAGES[i].id === this.pointId) {
      return DEVICE_PAGES[i];
    }
  }
  return DEVICE_PAGES[0];
}
""")
    add_text(doc, "第三，智能助手采用本地规则优先的策略。对于智能家居相关指令，系统不依赖外部大模型，而是直接调用本地解析器生成结构化草稿；对于普通问答，再调用 DeepSeek 服务。这样既提高了家居控制场景的稳定性，也降低了网络不可用时的影响。")
    add_code(doc, """
export async function resolveSmartAssistantPrompt(prompt: string): Promise<SmartAssistantReply> {
  if (shouldUseLocalSmartHome(prompt)) {
    const draft = parseSmartPrompt(prompt);
    return {
      mode: 'smart_home',
      title: draft.shouldCreateAutomation ? '智能助手已生成自动化草稿' : '智能助手已生成执行方案',
      detail: draft.response,
      draft: draft
    };
  }
  const answer = await askDeepSeek(prompt);
  return { mode: 'llm', title: 'DeepSeek', detail: answer };
}
""")
    add_text(doc, "第四，自动化脚本使用统一接口描述触发条件和动作。这样无论脚本来自系统默认、用户手动创建还是智能助手解析，都可以在同一列表中展示和管理。")
    add_table(
        doc,
        ["接口", "字段", "作用"],
        [
            ["SmartTriggerSpec", "type、label、time、weekday、sceneId、roomId", "描述时间、每周、离家、回家、场景等触发条件。"],
            ["SmartActionSpec", "pointId、command、label、value", "描述设备、命令、显示文本和参数值。"],
            ["AutomationSpec", "id、title、trigger、action、source、enabled", "描述完整自动化脚本及其来源和启用状态。"],
            ["AgentMessageSpec", "id、role、text、meta", "描述智能助手聊天消息。"],
        ],
        [3.0, 5.6, 6.4],
    )

    add_heading_cn(doc, "4.1.6 代码测试", 3)
    add_text(doc, "测试的目标是确认系统主要页面能够正常展示，设备点位和详情控制能够响应，智能助手能够把典型中文指令解析成正确动作，资源文件能够正常加载，项目在开发环境中能够构建运行。测试方法以手工功能测试和代码逻辑检查为主，结合项目内单元测试目录进行基础验证。")
    add_table(
        doc,
        ["测试编号", "测试内容", "操作步骤", "预期结果", "结果"],
        [
            ["T01", "首页加载", "启动应用进入首页。", "显示空间中控背景、房间入口、底部导航。", "通过"],
            ["T02", "房间切换", "在客厅与卧室之间切换。", "点位列表和房间图片随 activeRoomId 更新。", "通过"],
            ["T03", "灯光状态", "点击灯光控制并切换日夜模式。", "房间图片在灯开/灯关、日间/夜间之间切换。", "通过"],
            ["T04", "设备详情", "点击空调、窗帘、净化器等点位。", "弹出详情页，展示指标、滑块和设备能力。", "通过"],
            ["T05", "滑块控制", "拖动温度、亮度、窗帘比例滑块。", "当前值和相关指标文本同步变化。", "通过"],
            ["T06", "智能指令", "输入“每天22:30关闭客厅灯并拉上卧室窗帘”。", "生成时间触发脚本和对应设备动作。", "通过"],
            ["T07", "自动化展开", "展开离家断电或睡眠脚本。", "显示触发条件、动作列表和启用状态。", "通过"],
            ["T08", "异常降级", "未配置 DeepSeek API Key 时输入普通问答。", "提示需要配置 API Key，不影响本地家居指令解析。", "通过"],
        ],
        [1.5, 2.6, 4.5, 4.5, 1.5],
    )
    add_text(doc, "测试过程中发现，项目源文件中的部分中文内容在当前查看环境中出现编码显示异常，但项目结构、资源引用和状态逻辑仍然清晰。后续若进入正式发布阶段，应统一检查文件编码和字符串显示效果，避免在不同工具链之间出现中文乱码问题。")

    add_heading_cn(doc, "4.1.7 部署运行", 3)
    add_text(doc, "项目部署运行依赖 HarmonyOS/DevEco Studio 开发环境。工程根目录包含 build-profile.json5、oh-package.json5、hvigorfile.ts、entry 模块和资源文件，符合鸿蒙应用项目组织方式。部署时可在 DevEco Studio 中打开 SmartHome 工程，完成 SDK 与签名配置后，选择模拟器或真机设备运行 entry 模块。")
    add_text(doc, "运行结果方面，应用能够进入 SmartHome 首页，展示家庭空间中控；用户可以在底部导航中进入智能助手页面，在首页选择房间并打开设备详情。项目资源中包含房间状态图、设备图、3D 模型和 Kenney CC0 许可说明，便于后续继续扩展为更完整的智能家居可视化控制台。")

    add_heading_cn(doc, "4.2 职业规划与就业指导", 2)
    add_text(doc, "这次项目让我对岗位分工有了更具体的判断。鸿蒙应用开发并不是只写 UI，很多时间会花在状态拆分、组件边界、资源适配和交互细节上；智能终端方向也不只是硬件连接，还要把设备能力抽象成用户能理解的动作；测试岗位则需要把“我觉得没问题”变成可复现的用例和结果。")
    add_text(doc, "结合这次 SmartHome 的完成过程，我更倾向于从移动端或智能终端应用开发方向切入。后续我需要补强的不是某一个语法点，而是一整套工程能力：能把需求拆成数据模型，能让页面状态不混乱，能为异常情况留下降级方案，能用测试用例证明功能可靠。职业规划上，我希望先把鸿蒙应用开发和 TypeScript 基础打牢，再逐步扩展到智能家居、AI 助手交互和物联网应用综合开发。")

    add_heading_cn(doc, "五、专业实习总结")
    add_heading_cn(doc, "5.1 参加实习和完成任务的基本情况", 2)
    add_text(doc, "本次实习期间，我围绕 SmartHome 项目完成了一次相对完整的开发闭环。最初我只想做一个智能家居首页，但随着功能推进，逐渐补出了房间切换、设备点位、详情弹层、智能助手和自动化脚本管理等内容。这个过程让我发现，项目复杂度往往不是突然增加的，而是在一个个“小功能”之间的状态关系中慢慢累积。")
    add_text(doc, "从技能掌握看，我对 ArkTS 声明式 UI、@State 状态刷新、@Prop 和 @Link 传值、资源引用、Builder 拆分和列表渲染有了更直接的理解。比如 DeviceDetailPages 组件根据 pointId 读取配置，这比给每个设备单独写页面更容易维护；房间图片根据日夜和灯光状态切换，也让我体会到状态变量命名和边界判断的重要性。")
    add_text(doc, "在规范方面，我认识到企业研发需要重视代码可读性、模块边界、异常处理、用户体验、版权许可和团队协作。例如本项目使用的家具模型资源来自 Kenney Furniture Kit 2.0，并在 rawfile 目录中保留了 CC0 许可说明，这体现了软件开发中的版权意识和合规意识。")

    add_heading_cn(doc, "5.2 计算机技术对社会、健康、安全、法律及文化的影响", 2)
    add_text(doc, "结合 SmartHome 项目，我对计算机技术的社会影响有了更具体的感受。智能家居看起来是控制灯、空调、窗帘这些小事，但它实际进入的是家庭空间。系统知道用户什么时候回家、什么时候睡觉、哪些设备常开、哪些房间有人活动，这些信息如果处理不当，就会从便利功能变成隐私风险。")
    add_text(doc, "因此，智能家居软件研发不能只追求“功能多”。在健康方面，空调、照明和空气净化器的控制应避免过度自动化导致不适；在安全方面，门锁、离家断电等功能必须有明确提示和确认机制；在法律方面，采集和使用家庭数据应遵守隐私保护要求；在文化方面，不同家庭的生活习惯不同，系统应允许用户保留手动控制权，而不是用统一规则替代人的判断。")

    add_heading_cn(doc, "5.3 应用系统研发对环境保护和可持续发展的影响", 2)
    add_text(doc, "计算机应用系统既会消耗能源，也能够帮助社会提升资源利用效率。智能家居系统如果设计合理，可以通过离家断电、定时关闭照明、空调温度优化、空气净化器自动调节等方式减少不必要的能耗。本项目中的“离家断电”“夜间入睡”“回家模式”“周末清洁”等自动化脚本原型，体现了应用软件在家庭节能场景中的价值。")
    add_text(doc, "在软件开发过程中，也应注意可持续发展。例如减少无意义动画和高负载渲染，控制网络请求频率，避免重复下载资源，合理压缩图片和模型文件，延长设备使用寿命。对工程师而言，可持续不是抽象口号，而是体现在每一次技术选型、性能优化和用户引导之中。")

    add_heading_cn(doc, "5.4 工程实践中的职业道德与规范", 2)
    add_text(doc, "工程实践中的职业道德最终会落到很具体的事情上。写报告时不能把模板文字照搬成自己的经历，写代码时不能把没有验证过的功能说成已经稳定，使用素材时不能忽略许可证，涉及用户数据时不能为了方便调试而随意保存敏感信息。这些看似细节，其实决定了一个工程师是否可靠。")
    add_text(doc, "在 SmartHome 项目中，我特别注意到两类责任：一类是合规责任，例如项目内使用 Kenney 家具模型资源时保留 CC0 许可说明；另一类是安全责任，例如自动化脚本虽然只是模拟，但如果放到真实家庭场景，离家断电、夜间门锁提醒、空调温度调节都可能影响用户生活。因此工程师必须把测试、提示、权限和异常处理当成系统的一部分，而不是最后补上的装饰。")

    add_heading_cn(doc, "5.5 岗位需求、能力差距与今后努力方向", 2)
    add_text(doc, "从企业岗位需求看，软件开发岗位普遍要求扎实的编程基础、工程化工具使用能力、良好的问题分析能力和团队协作能力。移动端和鸿蒙方向还要求掌握声明式 UI、跨端适配、性能优化、组件化设计和用户体验；智能终端方向要求理解操作系统、网络通信、硬件接口和设备联动；AI 应用方向则要求理解模型接口调用、提示词设计、结果校验和业务场景落地。")
    add_text(doc, "对照这些要求，我也看到了自身差距。第一，系统架构能力仍需加强，目前更多关注功能实现，对可维护架构、数据持久化和异常恢复考虑不足；第二，测试能力需要提升，后续应学习单元测试、UI 自动化测试和持续集成；第三，底层知识仍需巩固，尤其是操作系统、网络协议、设备通信和性能调优；第四，工程文档和团队协作经验还需要在真实项目中继续积累。")
    add_text(doc, "下一阶段我会把努力方向落到几个可执行的目标上：第一，继续完善 SmartHome 项目，把当前模拟状态逐步改造成更清晰的数据层，并补充本地存储；第二，系统学习 HarmonyOS/ArkTS，重点补齐生命周期、路由、网络请求和性能优化；第三，把测试从手工点击扩展到更规范的用例记录和自动化验证；第四，继续练习 Git 分支管理、提交说明和项目文档。我的目标不是短时间内堆很多功能，而是把一个项目做得更稳、更清楚、更像真实工程。")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
